import io
import pytest
from fastapi.testclient import TestClient
import fitz  # PyMuPDF
import docx

from app.main import app
from app.services.parser import ResumeParser, ResumeParsingError
from app.services.ats_scorer import ATSScorer
from app.services.job_matcher import JobMatcher
from app.services.skills_extractor import SkillsExtractor
from app.services.semantic_matcher import SemanticMatcher
from app.services.ollama_service import OllamaService
from app.services.report_generator import ReportGenerator

client = TestClient(app)

# --- Synthetic Resume Generation Utilities ---
def create_synthetic_pdf(
    name: str = "Alex Morgan",
    email: str = "alex.morgan@example.com",
    phone: str = "+1-555-019-2834",
    skills: list = None,
    experience_bullets: list = None,
    education_text: str = "Bachelor of Science in Computer Science, Tech University, 2023"
) -> bytes:
    if skills is None:
        skills = ["Python", "FastAPI", "React", "PostgreSQL", "Docker", "Git", "RESTful APIs"]
    if experience_bullets is None:
        experience_bullets = [
            "Developed scalable microservices using Python and FastAPI, increasing API throughput by 40%.",
            "Engineered responsive user interfaces with React and TypeScript for 10,000+ active users.",
            "Architected PostgreSQL relational database schemas and optimized slow SQL queries.",
            "Automated continuous integration and deployment pipelines using Docker and GitHub Actions."
        ]

    doc = fitz.open()
    page = doc.new_page()
    
    y = 50
    page.insert_text((50, y), name, fontsize=18)
    y += 20
    page.insert_text((50, y), f"{email} | {phone} | linkedin.com/in/alexmorgan | github.com/alexmorgan", fontsize=9)
    y += 30
    page.insert_text((50, y), "TECHNICAL SKILLS", fontsize=12)
    y += 15
    page.insert_text((50, y), ", ".join(skills), fontsize=10)
    y += 30
    page.insert_text((50, y), "WORK EXPERIENCE", fontsize=12)
    y += 18
    for b in experience_bullets:
        page.insert_text((60, y), f"• {b}", fontsize=9)
        y += 15
    y += 20
    page.insert_text((50, y), "EDUCATION", fontsize=12)
    y += 15
    page.insert_text((60, y), f"• {education_text}", fontsize=9)
    y += 25
    page.insert_text((50, y), "PROJECTS", fontsize=12)
    y += 15
    page.insert_text((60, y), "• ResumeIQ: Full-stack ATS analyzer built with React, FastAPI, and Docker.", fontsize=9)

    pdf_bytes = doc.tobytes()
    doc.close()
    return pdf_bytes

def create_synthetic_docx(
    name: str = "Taylor Swiftly",
    email: str = "taylor.s@example.com",
    phone: "str" = "+1-555-321-7654"
) -> bytes:
    doc = docx.Document()
    doc.add_heading(name, level=0)
    doc.add_paragraph(f"{email} | {phone}")
    
    doc.add_heading("Technical Skills", level=1)
    doc.add_paragraph("Python, React, TypeScript, Docker, SQL, Git, RESTful APIs, Tailwind CSS")
    
    doc.add_heading("Work Experience", level=1)
    p = doc.add_paragraph()
    p.add_run("Software Engineer - CloudCorp (2022 - Present)\n")
    p.add_run("• Architected scalable web services using FastAPI and Python, improving uptime to 99.9%.\n")
    p.add_run("• Developed modern frontend components utilizing React, TypeScript, and Tailwind CSS.\n")
    p.add_run("• Deployed containerized microservices to cloud environments with Docker.")
    
    doc.add_heading("Education", level=1)
    doc.add_paragraph("B.S. in Computer Science - State University (2018 - 2022)")
    
    doc.add_heading("Projects", level=1)
    doc.add_paragraph("• Distributed Task Runner: Built with Python, Redis, and Docker.")

    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()

SAMPLE_JOB_DESCRIPTION = """
Senior Full Stack Engineer (React & Python)

Requirements:
- 3+ years of professional software development experience.
- Strong proficiency in Python, FastAPI, and RESTful APIs.
- Deep expertise in React, TypeScript, JavaScript, and modern CSS.
- Hands-on experience with SQL, PostgreSQL, and database design.
- Practical experience with Docker, Git version control, and CI/CD pipelines.
- Bachelor's degree in Computer Science or related quantitative field.

Preferred Qualifications:
- Familiarity with Kubernetes, Redis, and Cloud platforms (AWS/GCP).
- Experience with Unit Testing (pytest, jest).
- Excellent team collaboration and problem-solving skills.
"""

# --- 1. Health & AI Status Tests ---
def test_health_endpoint():
    response = client.get("/api/health")
    assert response.status_code == 200
    data = response.json()
    assert data["status"] == "healthy"
    assert "version" in data
    assert "ResumeIQ" in data["app_name"]

def test_ai_status_endpoint():
    response = client.get("/api/ai/status")
    assert response.status_code == 200
    data = response.json()
    assert "ollama_running" in data
    assert "current_model" in data
    assert isinstance(data["available_models"], list)

# --- 2. Parser Unit Tests ---
def test_parse_valid_pdf():
    pdf_bytes = create_synthetic_pdf()
    text, sections = ResumeParser.parse_document("test_resume.pdf", pdf_bytes)
    assert len(text) > 100
    assert "Alex Morgan" in sections.candidate_name
    assert sections.email == "alex.morgan@example.com"
    assert len(sections.skills) > 0
    assert len(sections.work_experience) > 0
    assert len(sections.education) > 0

def test_parse_valid_docx():
    docx_bytes = create_synthetic_docx()
    text, sections = ResumeParser.parse_document("test_resume.docx", docx_bytes)
    assert len(text) > 100
    assert "Taylor" in sections.candidate_name
    assert sections.email == "taylor.s@example.com"
    assert len(sections.skills) > 0
    assert len(sections.work_experience) > 0

def test_unsupported_file_extension():
    with pytest.raises(ResumeParsingError) as exc_info:
        ResumeParser.validate_file("resume.txt", b"plain text")
    assert "Unsupported file format" in str(exc_info.value)

def test_empty_file():
    with pytest.raises(ResumeParsingError) as exc_info:
        ResumeParser.validate_file("resume.pdf", b"")
    assert "empty" in str(exc_info.value)

def test_too_short_text():
    # PDF with only 5 words
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 50), "Only five short words here.")
    short_pdf = doc.tobytes()
    doc.close()

    with pytest.raises(ResumeParsingError) as exc_info:
        ResumeParser.parse_document("short.pdf", short_pdf)
    assert "insufficient or empty" in str(exc_info.value)

# --- 3. Deterministic ATS Scoring Tests ---
def test_ats_score_deterministic_and_range():
    pdf_bytes = create_synthetic_pdf()
    text, sections = ResumeParser.parse_document("alex.pdf", pdf_bytes)
    
    result1 = ATSScorer.evaluate(text, sections, SAMPLE_JOB_DESCRIPTION)
    result2 = ATSScorer.evaluate(text, sections, SAMPLE_JOB_DESCRIPTION)

    # Must be deterministic
    assert result1.total_score == result2.total_score
    # Score in [0, 100]
    assert 0 <= result1.total_score <= 100
    # Alex's resume has complete sections and contact info, should score >= 75
    assert result1.total_score >= 75
    assert len(result1.factors) == 5
    assert result1.status_label in ["Excellent", "Good", "Needs Improvement", "Critical Issues"]

def test_ats_score_handles_missing_sections_gracefully():
    # Incomplete resume without education or contact, but with enough words (> 25)
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text(
        (50, 50),
        "WORK EXPERIENCE\n"
        "• Developed distributed cloud software using Python and Docker for two years.\n"
        "• Engineered reliable backend systems, optimized database queries, and deployed microservices.\n"
        "• Collaborated with engineering teams across multiple agile sprint lifecycles.",
        fontsize=10
    )
    pdf_bytes = doc.tobytes()
    doc.close()

    text, sections = ResumeParser.parse_document("incomplete.pdf", pdf_bytes)
    result = ATSScorer.evaluate(text, sections, SAMPLE_JOB_DESCRIPTION)
    
    assert 0 <= result.total_score <= 100
    # Should flag missing contact and education in warnings
    assert any("contact" in w.lower() for w in result.formatting_warnings)
    assert any("education" in w.lower() for w in result.formatting_warnings)

# --- 4. Job Matching Engine Tests ---
def test_job_match_score_range_and_breakdown():
    pdf_bytes = create_synthetic_pdf()
    text, sections = ResumeParser.parse_document("alex.pdf", pdf_bytes)
    
    match_result = JobMatcher.evaluate(text, sections, SAMPLE_JOB_DESCRIPTION)
    assert 0 <= match_result.total_score <= 100
    assert match_result.breakdown.required_skills_score >= 60
    assert len(match_result.matched_skills) > 0
    assert match_result.semantic_similarity_percent >= 0

def test_skills_extractor_matches_and_partials():
    resume_text = "Experienced in Python, React, and PostgreSQL. Familiar with Docker and SQL."
    jd_text = """
    Requirements:
    - Must have Python and React.
    - PostgreSQL database experience.
    - Kubernetes container management.
    - C++ programming.
    """
    matched, partial, missing = SkillsExtractor.match_skills(resume_text, jd_text)
    matched_names = [s.skill for s in matched]
    partial_names = [s.skill for s in partial]
    missing_names = [s.skill for s in missing]

    # Python, React, PostgreSQL should be matched
    assert "Python" in matched_names
    assert "React" in matched_names
    assert "PostgreSQL" in matched_names

    # Kubernetes should be partially matched via Docker!
    assert "Kubernetes" in partial_names

    # C++ should be missing
    assert "C++" in missing_names

# --- 5. Semantic Matcher & AI Fallback Tests ---
def test_semantic_matcher_fallback():
    score, engine = SemanticMatcher.compute_similarity(
        "Software Engineer with Python and React experience",
        "Looking for a Python and React developer"
    )
    assert 0 <= score <= 100
    assert len(engine) > 0

def test_ollama_fallback_engine():
    pdf_bytes = create_synthetic_pdf()
    text, sections = ResumeParser.parse_document("alex.pdf", pdf_bytes)
    ats = ATSScorer.evaluate(text, sections, SAMPLE_JOB_DESCRIPTION)
    match = JobMatcher.evaluate(text, sections, SAMPLE_JOB_DESCRIPTION)

    # Force rule-based fallback
    explanation = OllamaService._generate_rule_based_explanation(sections, ats, match)
    assert explanation.is_fallback is True
    assert explanation.ai_provider == "rule_based"
    assert len(explanation.strengths) > 0
    assert len(explanation.weaknesses) > 0
    assert len(explanation.improvement_recommendations) > 0
    assert len(explanation.suggested_bullet_rewrites) > 0

# --- 6. Full End-to-End API Analysis Test ---
def test_full_analyze_endpoint():
    pdf_bytes = create_synthetic_pdf()
    response = client.post(
        "/api/analyze",
        files={"file": ("alex_resume.pdf", pdf_bytes, "application/pdf")},
        data={"job_description": SAMPLE_JOB_DESCRIPTION}
    )
    assert response.status_code == 200
    data = response.json()
    assert "id" in data
    assert data["ats_result"]["total_score"] >= 60
    assert data["job_match_result"]["total_score"] >= 60
    assert len(data["job_match_result"]["matched_skills"]) > 0
    assert len(data["ai_explanation"]["strengths"]) > 0

    analysis_id = data["id"]

    # Test history retrieval
    hist_resp = client.get("/api/analyses")
    assert hist_resp.status_code == 200
    hist = hist_resp.json()
    assert any(h["id"] == analysis_id for h in hist)

    # Test single analysis get
    single_resp = client.get(f"/api/analyses/{analysis_id}")
    assert single_resp.status_code == 200
    assert single_resp.json()["id"] == analysis_id

    # Test printable HTML report
    report_resp = client.get(f"/api/analyses/{analysis_id}/report")
    assert report_resp.status_code == 200
    assert "text/html" in report_resp.headers["content-type"]
    assert "ResumeIQ Analysis Report" in report_resp.text
    assert "ATS READINESS SCORE" in report_resp.text

    # Test delete
    del_resp = client.delete(f"/api/analyses/{analysis_id}")
    assert del_resp.status_code == 200
