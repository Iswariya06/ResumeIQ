import os
import re
from typing import Dict, List, Tuple, Optional
import fitz  # PyMuPDF
import docx
from app.core.config import settings
from app.schemas.analysis import ExtractedSections

EMAIL_REGEX = re.compile(r'[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+')
PHONE_REGEX = re.compile(r'(\+?\d{1,3}[-.\s]?)?(\(?\d{3}\)?[-.\s]?)?\d{3}[-.\s]?\d{4}')

COMMON_SECTION_PATTERNS = {
    "experience": [
        r"^work\s+experience", r"^professional\s+experience", r"^experience",
        r"^employment\s+history", r"^work\s+history", r"^internships?"
    ],
    "education": [
        r"^education", r"^academic\s+background", r"^academics?", r"^qualifications?"
    ],
    "skills": [
        r"^technical\s+skills", r"^skills?\s+(&|and)?\s+technologies", r"^skills?",
        r"^core\s+competencies", r"^technologies", r"^proficiencies"
    ],
    "projects": [
        r"^projects?", r"^key\s+projects?", r"^academic\s+projects?", r"^personal\s+projects?"
    ],
    "certifications": [
        r"^certifications?", r"^licenses?\s+(&|and)?\s+certifications?", r"^certificates?"
    ],
    "achievements": [
        r"^achievements?", r"^awards?\s+(&|and)?\s+honors?", r"^honors?\s+(&|and)?\s+awards?",
        r"^accomplishments?"
    ]
}

class ResumeParsingError(Exception):
    pass

class ResumeParser:
    @staticmethod
    def validate_file(filename: str, file_bytes: bytes) -> str:
        """Validate file size and extension."""
        ext = os.path.splitext(filename)[1].lower()
        if ext not in settings.ALLOWED_EXTENSIONS:
            raise ResumeParsingError(
                f"Unsupported file format '{ext}'. Only PDF (.pdf) and Word (.docx) documents are supported."
            )
        
        if len(file_bytes) == 0:
            raise ResumeParsingError("The uploaded file is empty (0 bytes).")
            
        if len(file_bytes) > settings.MAX_UPLOAD_SIZE_BYTES:
            max_mb = settings.MAX_UPLOAD_SIZE_BYTES / (1024 * 1024)
            raise ResumeParsingError(
                f"File size exceeds the {max_mb:.0f}MB limit. Please upload a smaller file."
            )
            
        return ext

    @classmethod
    def extract_text_from_pdf(cls, file_bytes: bytes) -> str:
        """Extract text from PDF using PyMuPDF (fitz)."""
        try:
            doc = fitz.open(stream=file_bytes, filetype="pdf")
        except Exception as e:
            raise ResumeParsingError(f"Unable to read PDF file. The file may be corrupt or invalid: {str(e)}")

        if doc.is_encrypted:
            doc.close()
            raise ResumeParsingError(
                "The PDF is encrypted or password-protected. Please provide an unencrypted resume."
            )

        text_parts = []
        for page_idx in range(len(doc)):
            page = doc[page_idx]
            text = page.get_text("text")
            if text:
                text_parts.append(text)

        doc.close()
        full_text = "\n".join(text_parts).strip()
        return full_text

    @classmethod
    def extract_text_from_docx(cls, file_bytes: bytes, temp_path: str) -> str:
        """Extract text from DOCX using python-docx."""
        try:
            with open(temp_path, "wb") as f:
                f.write(file_bytes)
            
            doc = docx.Document(temp_path)
            text_parts = []
            
            # Paragraphs
            for p in doc.paragraphs:
                if p.text.strip():
                    text_parts.append(p.text.strip())
            
            # Tables
            for table in doc.tables:
                for row in table.rows:
                    row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_texts:
                        text_parts.append(" | ".join(row_texts))
                        
            return "\n".join(text_parts).strip()
        except Exception as e:
            raise ResumeParsingError(f"Unable to read DOCX file. The document may be corrupted: {str(e)}")
        finally:
            if os.path.exists(temp_path):
                try:
                    os.remove(temp_path)
                except Exception:
                    pass

    @classmethod
    def parse_document(cls, filename: str, file_bytes: bytes, temp_dir: str = ".") -> Tuple[str, ExtractedSections]:
        """Validates, extracts raw text, and segments the resume into structured sections."""
        ext = cls.validate_file(filename, file_bytes)
        
        if ext == ".pdf":
            raw_text = cls.extract_text_from_pdf(file_bytes)
        elif ext == ".docx":
            temp_path = os.path.join(temp_dir, f"temp_{os.getpid()}_{filename}")
            raw_text = cls.extract_text_from_docx(file_bytes, temp_path)
        else:
            raise ResumeParsingError(f"Unsupported extension {ext}")

        words = raw_text.split()
        if len(words) < 25:
            raise ResumeParsingError(
                "Extracted text is insufficient or empty (< 25 words). "
                "If this is a scanned image PDF, please supply a searchable text-based PDF or DOCX file."
            )

        sections = cls.extract_sections(raw_text)
        return raw_text, sections

    @classmethod
    def extract_sections(cls, raw_text: str) -> ExtractedSections:
        """Segment raw text into contact info and structured resume sections."""
        lines = [line.strip() for line in raw_text.splitlines() if line.strip()]
        
        # 1. Candidate Name detection
        candidate_name = cls._detect_candidate_name(lines)
        
        # 2. Email and Phone detection
        email = None
        phone = None
        email_match = EMAIL_REGEX.search(raw_text)
        if email_match:
            email = email_match.group(0)
            
        phone_match = PHONE_REGEX.search(raw_text)
        if phone_match and len(phone_match.group(0).strip()) >= 7:
            phone = phone_match.group(0).strip()

        # 3. Section grouping
        section_buckets: Dict[str, List[str]] = {
            "experience": [],
            "education": [],
            "skills": [],
            "projects": [],
            "certifications": [],
            "achievements": [],
            "other": []
        }
        
        detected_headers: List[str] = []
        current_section = "other"

        for line in lines:
            normalized_line = line.lower().strip()
            # Remove trailing colons or dashes
            normalized_line_clean = re.sub(r'[:\-_#*]+$', '', normalized_line).strip()
            
            # Check if this line matches a standard section header
            matched_sec = None
            if len(normalized_line_clean.split()) <= 4 and len(normalized_line_clean) < 40:
                for sec_key, patterns in COMMON_SECTION_PATTERNS.items():
                    for pattern in patterns:
                        if re.match(pattern, normalized_line_clean):
                            matched_sec = sec_key
                            break
                    if matched_sec:
                        break

            if matched_sec:
                current_section = matched_sec
                detected_headers.append(line)
            else:
                section_buckets[current_section].append(line)

        return ExtractedSections(
            candidate_name=candidate_name,
            email=email,
            phone=phone,
            education=section_buckets["education"],
            skills=section_buckets["skills"],
            work_experience=section_buckets["experience"],
            projects=section_buckets["projects"],
            certifications=section_buckets["certifications"],
            achievements=section_buckets["achievements"],
            raw_word_count=len(raw_text.split()),
            section_headers_detected=detected_headers
        )

    @staticmethod
    def _detect_candidate_name(lines: List[str]) -> str:
        """Infer candidate name from header lines."""
        for line in lines[:6]:
            # Skip emails, phone numbers, links, common headers
            if "@" in line or "http" in line or "linkedin" in line or "github" in line:
                continue
            cleaned = re.sub(r'[^a-zA-Z\s\.]', '', line).strip()
            words = cleaned.split()
            if 2 <= len(words) <= 4 and len(cleaned) < 40:
                # Disqualify if it contains words like 'resume', 'curriculum', 'developer', etc.
                lower = cleaned.lower()
                if any(w in lower for w in ["resume", "curriculum", "vitae", "profile", "contact", "summary", "engineer", "developer"]):
                    continue
                return cleaned
        return "Candidate"
